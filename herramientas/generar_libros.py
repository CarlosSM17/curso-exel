"""Genera los libros de Excel (y la carta de Word) del curso, versión 2: 13 módulos.

  datos/la-taza-base.xlsx                 libro de datos del caso
  laboratorios/labXX-*.xlsx               13 laboratorios (etapa 4) y lab13-carta.docx
  integradores/int-*.xlsx                 retos integradores
  <privado>/soluciones/laboratorios/      soluciones de los laboratorios   } fuera del sitio:
  <privado>/soluciones/integradores/      soluciones de los integradores   } no se publican

<privado> es la carpeta hermana «curso-excel-la-taza-v2-privado» (o la variable de
entorno LT_PRIVADO). Después ejecuta herramientas/proteger_docente.py para cifrar las
soluciones dentro del área docente.

Uso:  python herramientas/generar_libros.py
La disposición de las celdas coincide con las pruebas de assets/js/modulos/mXX.js
y assets/js/paginas.js: si cambias una, cambia la otra.
"""
import datetime as dt
import os
import pathlib
import random
import re
import zipfile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.worksheet.scenario import InputCells, Scenario, ScenarioList
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

import lataza_datos as LD

RAIZ = pathlib.Path(__file__).resolve().parent.parent
PRIVADO = pathlib.Path(os.environ.get("LT_PRIVADO") or RAIZ.parent / (RAIZ.name + "-privado"))
D = LD.generar()
PROD = {p["codigo"]: p for p in D["productos"]}
CATS = D["categorias"]
MESES = D["mesesCortos"]
DIAS = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]

NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=16, color="5B3A29")
RELLENO_ENC = PatternFill("solid", fgColor="F3EBDF")
RELLENO_ENTRADA = PatternFill("solid", fgColor="FFF8D6")
FINO = Side(style="thin", color="555555")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
ROJO = PatternFill("solid", fgColor="FFC7CE")
AMARILLO = PatternFill("solid", fgColor="FFEB9C")
MONEDA = "$#,##0.00"
MONEDA0 = "$#,##0"


# ------------------------------------------------------------------ utilidades
def serial(iso):
    y, m, d = map(int, iso.split("-"))
    return (dt.date(y, m, d) - dt.date(1899, 12, 30)).days


def poner_fecha(celda, iso):
    celda.value = serial(iso)
    celda.number_format = "dd/mm/yyyy"


def encabezados(ws, fila, valores, col=1, relleno=True):
    for j, v in enumerate(valores):
        c = ws.cell(row=fila, column=col + j, value=v)
        c.font = NEGRITA
        if relleno:
            c.fill = RELLENO_ENC


def anchos(ws, **kw):
    for col, w in kw.items():
        ws.column_dimensions[col].width = w


def libro(titulo, contexto, pasos, criterios=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Instrucciones"
    ws["A1"] = titulo
    ws["A1"].font = TITULO
    ws["A3"] = contexto
    ws["A5"] = "Lo que debes hacer"
    ws["A5"].font = NEGRITA
    fila = 6
    for i, p in enumerate(pasos, 1):
        ws.cell(row=fila, column=1, value=f"{i}. {p}")
        fila += 1
    if criterios:
        fila += 1
        ws.cell(row=fila, column=1, value="El verificador revisará").font = NEGRITA
        fila += 1
        for c in criterios:
            ws.cell(row=fila, column=1, value="• " + c)
            fila += 1
    fila += 1
    ws.cell(row=fila, column=1, value="Importante: no cambies los nombres de las hojas ni muevas los datos de lugar. "
            "Al terminar, guarda el libro y súbelo en la página del módulo (etapa 4).").font = Font(italic=True, color="7A4D0E")
    ws.column_dimensions["A"].width = 120
    for r in ws.iter_rows(min_row=1, max_row=fila):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return wb


def hoja_solucion(wb, lineas):
    ws = wb.create_sheet("Solución")
    ws["A1"] = "Solución paso a paso (para el docente)"
    ws["A1"].font = TITULO
    for i, l in enumerate(lineas, 3):
        ws.cell(row=i, column=1, value=l)
    ws.column_dimensions["A"].width = 120
    return ws


def bordes(ws, rango):
    for fila in ws[rango]:
        for c in fila:
            c.border = BORDE


def guardar(wb, ruta):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def numerar(wb, nombre):
    """Los constructores de la versión 1 se reutilizan: ajusta «Laboratorio N» al número nuevo."""
    n = int(nombre[3:5])
    a1 = wb["Instrucciones"]["A1"]
    a1.value = re.sub(r"^Laboratorio \d+", f"Laboratorio {n}", a1.value)
    return wb


def lab(nombre, constructor):
    guardar(numerar(constructor(False), nombre), RAIZ / "laboratorios" / f"{nombre}.xlsx")
    guardar(numerar(constructor(True), nombre), PRIVADO / "soluciones" / "laboratorios" / f"{nombre}-solucion.xlsx")


def integrador(nombre, constructor):
    guardar(constructor(False), RAIZ / "integradores" / f"{nombre}.xlsx")
    guardar(constructor(True), PRIVADO / "soluciones" / "integradores" / f"{nombre}-solucion.xlsx")


def hoja_ventas(wb, nombre="Ventas"):
    ws = wb.create_sheet(nombre)
    encabezados(ws, 1, ["Folio", "Fecha", "Sucursal", "Código", "Producto", "Categoría", "Cantidad", "Precio", "Total", "Vendedor", "Pago"])
    for i, v in enumerate(D["ventas"], 2):
        ws.cell(row=i, column=1, value=v["folio"])
        poner_fecha(ws.cell(row=i, column=2), v["fecha"])
        for j, k in enumerate(["sucursal", "codigo", "producto", "categoria", "cantidad", "precio", "total", "vendedor", "pago"], 3):
            ws.cell(row=i, column=j, value=v[k])
        ws.cell(row=i, column=8).number_format = MONEDA0
        ws.cell(row=i, column=9).number_format = MONEDA0
    anchos(ws, A=9, B=11, C=10, D=10, E=22, F=15, G=10, H=9, I=10, J=18, K=14)
    ws.freeze_panes = "A2"
    return ws


def hoja_productos(wb):
    ws = wb.create_sheet("Productos")
    encabezados(ws, 1, ["Código", "Producto", "Categoría", "Precio", "Costo"])
    for i, p in enumerate(D["productos"], 2):
        for j, k in enumerate(["codigo", "nombre", "categoria", "precio", "costo"], 1):
            ws.cell(row=i, column=j, value=p[k])
        ws.cell(row=i, column=4).number_format = MONEDA0
        ws.cell(row=i, column=5).number_format = MONEDA0
    anchos(ws, A=10, B=22, C=15, D=9, E=9)
    return ws


# ============================================================ libro de datos
def libro_base():
    wb = Workbook()
    ws = wb.active
    ws.title = "Léeme"
    ws["A1"] = "Cafetería La Taza · datos del caso"
    ws["A1"].font = TITULO
    filas = [
        "Libro de datos del curso «Excel con La Taza». Todos los módulos usan estas mismas tablas.",
        "Productos: catálogo de 20 productos con precio y costo.",
        "Ventas: 720 registros de 2025 (día, sucursal y producto).",
        "Empleados: 20 vendedores con sus ventas de octubre y noviembre.",
        "Clientes: base del programa de lealtad (sin limpiar).",
        "Pedidos: 40 pedidos de catering con fecha de compromiso y estado.",
        "Inventario y Presupuesto: insumos y gastos del semestre.",
    ]
    for i, t in enumerate(filas, 3):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 100
    hoja_productos(wb)
    hoja_ventas(wb)
    ws = wb.create_sheet("Empleados")
    encabezados(ws, 1, ["Nombre", "Sucursal", "Puesto", "Ingreso", "Ventas octubre", "Ventas noviembre"])
    for i, e in enumerate(D["empleados"], 2):
        ws.cell(row=i, column=1, value=e["nombre"])
        ws.cell(row=i, column=2, value=e["sucursal"])
        ws.cell(row=i, column=3, value=e["puesto"])
        poner_fecha(ws.cell(row=i, column=4), e["ingreso"])
        ws.cell(row=i, column=5, value=e["ventasMes"]).number_format = MONEDA0
        ws.cell(row=i, column=6, value=e["ventasMesLab"]).number_format = MONEDA0
    anchos(ws, A=20, B=10, C=12, D=12, E=15, F=17)
    ws = wb.create_sheet("Clientes")
    encabezados(ws, 1, ["ID", "Nombre completo", "Teléfono", "Correo", "Sucursal", "Registro", "Puntos"])
    for i, c in enumerate(D["clientesSucios"], 2):
        for j, k in enumerate(["id", "nombre", "telefono", "correo", "sucursal"], 1):
            ws.cell(row=i, column=j, value=int(c[k]) if k == "telefono" else c[k])
        poner_fecha(ws.cell(row=i, column=6), c["registro"])
        ws.cell(row=i, column=7, value=c["puntos"])
    anchos(ws, A=9, B=22, C=13, D=30, E=10, F=12, G=8)
    escribir_pedidos(wb.create_sheet("Pedidos"), D["pedidos"])
    ws = wb.create_sheet("Inventario")
    escribir_inventario(ws, D["inventario"])
    ws = wb.create_sheet("Presupuesto")
    escribir_presupuesto(ws)
    guardar(wb, RAIZ / "datos" / "la-taza-base.xlsx")


def escribir_pedidos(ws, pedidos, importe=False):
    cols = ["Folio", "Fecha", "Cliente", "Sucursal", "Producto", "Cantidad", "Precio", "Compromiso", "Estado"]
    encabezados(ws, 1, cols)
    for i, p in enumerate(pedidos, 2):
        ws.cell(row=i, column=1, value=p["folio"])
        poner_fecha(ws.cell(row=i, column=2), p["fecha"])
        ws.cell(row=i, column=3, value=p["cliente"])
        ws.cell(row=i, column=4, value=p["sucursal"])
        ws.cell(row=i, column=5, value=p["producto"])
        ws.cell(row=i, column=6, value=p["cantidad"])
        ws.cell(row=i, column=7, value=p["precio"]).number_format = MONEDA0
        poner_fecha(ws.cell(row=i, column=8), p["compromiso"])
        ws.cell(row=i, column=9, value=p["estado"])
    anchos(ws, A=8, B=11, C=24, D=10, E=22, F=10, G=9, H=12, I=11)


def escribir_inventario(ws, inv):
    encabezados(ws, 1, ["Insumo", "Unidad", "Existencia", "Mínimo", "Máximo", "Costo unitario", "Proveedor", "Cobertura"], relleno=False)
    for i, x in enumerate(inv, 2):
        for j, k in enumerate(["insumo", "unidad", "existencia", "minimo", "maximo", "costo", "proveedor"], 1):
            ws.cell(row=i, column=j, value=x[k])
        ws.cell(row=i, column=6).number_format = MONEDA
        ws.cell(row=i, column=8, value=f"=C{i}/D{i}").number_format = "0%"
    anchos(ws, A=24, B=9, C=11, D=9, E=9, F=14, G=18, H=11)


def escribir_presupuesto(ws):
    ws["A1"] = "Presupuesto de gastos · 1.er semestre 2026"
    ws["A1"].font = TITULO
    encabezados(ws, 3, ["Concepto", "Ene", "Feb", "Mar", "Abr", "May", "Jun", "Total"])
    for i, p in enumerate(D["presupuesto"], 4):
        ws.cell(row=i, column=1, value=p["concepto"])
        for j, v in enumerate(p["meses"], 2):
            ws.cell(row=i, column=j, value=v).number_format = MONEDA0
        ws.cell(row=i, column=8, value=f"=SUM(B{i}:G{i})").number_format = MONEDA0
    ws.cell(row=12, column=1, value="Total").font = NEGRITA
    for j in range(2, 9):
        L = get_column_letter(j)
        c = ws.cell(row=12, column=j, value=f"=SUM({L}4:{L}11)")
        c.number_format = MONEDA0
        c.font = NEGRITA
    anchos(ws, A=22)


# ================================================================ laboratorios
def lab01(sol):
    wb = libro("Laboratorio 1 · Lista de precios de té, postres y alimentos",
               "La segunda parte de la lista de precios tiene los mismos problemas que viste en el simulador.",
               ["Inserta el Brownie (POS-002, Postres, $40) entre el Pay de queso y la Galleta de avena.",
                "Completa los códigos que faltan con autorrelleno.",
                "Corrige «Galleta de avna» y el precio del Panini caprese ($85).",
                "Reemplaza la categoría «Postre» por «Postres» en toda la lista.",
                "Nivel Experto: sube 10 % los precios de Alimentos con Pegado especial › Multiplicar."],
               ["Códigos, productos, categorías y precios correctos en A2:D8.", "Precios como números."])
    ws = wb.create_sheet("Precios")
    encabezados(ws, 1, ["Código", "Producto", "Categoría", "Precio"])
    if sol:
        filas = [["TE-001", "Chai latte", "Té", 54], ["TE-002", "Té verde", "Té", 38], ["POS-001", "Pay de queso", "Postres", 55],
                 ["POS-002", "Brownie", "Postres", 40], ["POS-003", "Galleta de avena", "Postres", 25], ["ALI-001", "Sándwich de pavo", "Alimentos", 78],
                 ["ALI-002", "Panini caprese", "Alimentos", 85]]
    else:
        filas = [["TE-001", "Chai latte", "Té", 54], [None, "Té verde", "Té", 38], ["POS-001", "Pay de queso", "Postre", 55],
                 [None, "Galleta de avna", "Postre", 25], ["ALI-001", "Sándwich de pavo", "Alimentos", 78], [None, "Panini caprese", "Alimentos", 8.5]]
    for i, f in enumerate(filas, 2):
        for j, v in enumerate(f, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=4).number_format = MONEDA
    anchos(ws, A=10, B=22, C=14, D=10)
    if sol:
        hoja_solucion(wb, ["Esta solución corresponde al nivel Analista.",
                           "Nivel Experto: escribe 1.1 en una celda libre, cópiala, selecciona D7:D8 y usa Pegado especial › Valores + Multiplicar.",
                           "Resultado Experto: Sándwich de pavo $85.80 y Panini caprese $93.50 (valores, sin fórmulas); después borra la celda con 1.1."])
    return wb


def tickets_norte():
    rng = random.Random(202)
    cods = ["CAF-002", "CAF-004", "PAN-002", "CAF-003", "FRI-002", "PAN-001", "CAF-006", "ALI-002", "FRI-001", "POS-002", "CAF-001", "TE-001", "FRI-003", "CAF-004"]
    horas = ["07:05", "07:31", "07:52", "08:16", "08:58", "09:20", "09:44", "10:30", "11:15", "12:48", "13:36", "15:10", "17:05", "18:40"]
    return [{"ticket": 2101 + i, "hora": h, "producto": PROD[c]["nombre"], "cantidad": rng.randint(1, 4), "precio": PROD[c]["precio"]} for i, (c, h) in enumerate(zip(cods, horas))]


def lab02(sol):
    wb = libro("Laboratorio 2 · Corte de caja de la sucursal Norte", "Mismo corte que en el simulador, con los tickets de otro día.",
               ["Calcula el importe de cada ticket en la columna F (cantidad × precio).", "Completa I2:I6 con SUMA, PROMEDIO, MAX, MIN y CONTAR sobre la columna F.",
                "Calcula en I10 la diferencia de caja: dinero contado − (fondo inicial + total vendido).", "Nivel Experto: en I11 calcula el IVA incluido en las ventas (precios con 16 % de IVA)."])
    ws = wb.create_sheet("Corte")
    encabezados(ws, 1, ["Ticket", "Hora", "Producto", "Cantidad", "Precio", "Importe"])
    ws["H1"] = "Resumen del día"
    ws["H1"].font = Font(bold=True, size=13)
    t = tickets_norte()
    total = 0
    for i, x in enumerate(t, 2):
        ws.cell(row=i, column=1, value=x["ticket"])
        hh, mm = map(int, x["hora"].split(":"))
        c = ws.cell(row=i, column=2, value=dt.time(hh, mm))
        c.number_format = "hh:mm"
        ws.cell(row=i, column=3, value=x["producto"])
        ws.cell(row=i, column=4, value=x["cantidad"])
        ws.cell(row=i, column=5, value=x["precio"]).number_format = MONEDA
        ws.cell(row=i, column=6).number_format = MONEDA
        total += x["cantidad"] * x["precio"]
    etiquetas = {2: "Total vendido", 3: "Ticket promedio", 4: "Venta mayor", 5: "Venta menor", 6: "Número de tickets", 8: "Fondo inicial", 9: "Dinero contado al cierre", 10: "Diferencia de caja", 11: "IVA incluido en las ventas"}
    for r, txt in etiquetas.items():
        ws.cell(row=r, column=8, value=txt).font = NEGRITA
        if r != 6:
            ws.cell(row=r, column=9).number_format = MONEDA
    ws["I8"] = 1500
    ws["I9"] = 1500 + total - 35
    if sol:
        for i in range(2, 16):
            ws.cell(row=i, column=6, value=f"=D{i}*E{i}")
        ws["I2"], ws["I3"], ws["I4"], ws["I5"], ws["I6"] = "=SUM(F2:F15)", "=AVERAGE(F2:F15)", "=MAX(F2:F15)", "=MIN(F2:F15)", "=COUNT(F2:F15)"
        ws["I10"], ws["I11"] = "=I9-(I8+I2)", "=I2-I2/1.16"
    anchos(ws, A=8, B=8, C=22, D=10, E=10, F=12, G=3, H=26, I=13)
    return wb


def lab03(sol):
    wb = libro("Laboratorio 3 · Reporte semanal de la sucursal Norte", "El reporte tiene los números correctos pero no tiene formato.",
               ["Destaca el título (A1): negrita y tamaño 14 o más (o estilo Título).", "Da formato de fecha a B2 y D2.", "Encabezados A4:E4 en negrita, con relleno y centrados.",
                "Precio y Ventas (B5:B12, D5:D13) en moneda; Participación (E5:E13) en porcentaje.", "Bordes en A4:E13 y fila de total en negrita.",
                "Nivel Experto: estilos de celda Título (A1) y Total (fila 13); columna A más ancha."])
    ws = wb.create_sheet("Reporte")
    ws["A1"] = "Reporte semanal de ventas · Sucursal Norte"
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Del", serial("2025-10-13"), "al", serial("2025-10-19")
    for j, v in enumerate(["Producto", "Precio", "Unidades", "Ventas", "Participación"], 1):
        ws.cell(row=4, column=j, value=v)
    filas = D["semanaSucursal"]["Norte"]
    for i, p in enumerate(filas, 5):
        ws.cell(row=i, column=1, value=p["producto"])
        ws.cell(row=i, column=2, value=p["precio"])
        ws.cell(row=i, column=3, value=sum(p["cantidades"]))
        ws.cell(row=i, column=4, value=f"=B{i}*C{i}")
        ws.cell(row=i, column=5, value=f"=D{i}/$D$13")
    ws["A13"], ws["C13"], ws["D13"], ws["E13"] = "Total", "=SUM(C5:C12)", "=SUM(D5:D12)", "=SUM(E5:E12)"
    if sol:
        ws["A1"].style = "Title"
        for d in ("B2", "D2"):
            ws[d].number_format = "dd/mm/yyyy"
        bordes(ws, "A4:E13")
        for c in ws[4]:
            if c.column <= 5:
                c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, Alignment(horizontal="center")
        for c in ws[13]:
            if c.column <= 5:
                c.style = "Total"
        for r in range(5, 14):
            ws.cell(row=r, column=2).number_format = MONEDA
            ws.cell(row=r, column=4).number_format = MONEDA
            ws.cell(row=r, column=5).number_format = "0.0%"
        anchos(ws, A=22)
    return wb


def lab04(sol):
    wb = libro("Laboratorio 4 · Encabezado del menú", "Descarga también el logo (logo-la-taza.png) desde la página del módulo.",
               ["Inserta el logo en la parte superior.", "Inserta una forma con el texto «Menú de temporada».", "Inserta un icono de café.",
                "En A5 inserta un vínculo con el texto «Ver menú en línea» a https://lataza.mx/menu.", "Nivel Experto: texto alternativo en logo e icono; cuadro de texto con «Abierto de 7:00 a 21:00»."])
    ws = wb.create_sheet("Menú")
    encabezados(ws, 7, ["Bebida", "Categoría", "Precio"])
    beb = [p for p in D["productos"] if p["codigo"][:3] in ("CAF", "FRI") or p["codigo"].startswith("TE")]
    for i, p in enumerate(beb, 8):
        ws.cell(row=i, column=1, value=p["nombre"])
        ws.cell(row=i, column=2, value=p["categoria"])
        ws.cell(row=i, column=3, value=p["precio"]).number_format = MONEDA0
    anchos(ws, A=24, B=16, C=10)
    if sol:
        img = Image(str(RAIZ / "assets" / "img" / "logo-la-taza.png"))
        img.width = img.height = 90
        ws.add_image(img, "A1")
        ws["A5"] = "Ver menú en línea"
        ws["A5"].hyperlink = "https://lataza.mx/menu"
        ws["A5"].style = "Hyperlink"
        hoja_solucion(wb, ["1. Insertar › Imágenes › Este dispositivo › logo-la-taza.png.", "2. Insertar › Formas › Cinta; escribe «Menú de temporada».",
                           "3. Insertar › Iconos › busca «café».", "4. A5 › Insertar › Vínculo › texto «Ver menú en línea», dirección https://lataza.mx/menu.",
                           "Nota: openpyxl no crea formas ni iconos; esta solución incluye solo el logo y el vínculo."])
    return wb


def semana43(suc):
    rng = random.Random(4300 + len(suc))
    return [{"producto": p["producto"], "cantidades": [max(3, round(c * rng.uniform(0.85, 1.15))) for c in p["cantidades"]]} for p in D["semanaSucursal"][suc]]


def lab05(sol):
    wb = libro("Laboratorio 5 · Organizar el libro de la semana 43", "Hoja1, Hoja2 y Hoja3 son las sucursales (el nombre está en A1 de cada una). Hoja4 es el resumen.",
               ["Cambia el nombre de cada hoja de sucursal según su título en A1.", "Cambia Hoja4 a «Resumen» y colócala antes de las hojas de sucursal.",
                "Pon color a las pestañas de las sucursales.", "En Resumen!B2:B4 trae el total (I11) de cada sucursal con referencias a las otras hojas; en B5 súmalos.",
                "Nivel Experto: en Centro inmoviliza las filas 1–2 y la columna A."])
    orden = [("Hoja1", "Sur"), ("Hoja2", "Centro"), ("Hoja3", "Norte")]
    hojas = {}
    for nombre, suc in orden:
        ws = wb.create_sheet(suc if sol else nombre)
        hojas[suc] = ws
        ws["A1"] = f"Sucursal {suc} · Semana 43 (unidades vendidas)"
        ws["A1"].font = Font(bold=True, size=13)
        encabezados(ws, 2, ["Producto"] + DIAS + ["Total"])
        for i, p in enumerate(semana43(suc), 3):
            ws.cell(row=i, column=1, value=p["producto"])
            for j, v in enumerate(p["cantidades"], 2):
                ws.cell(row=i, column=j, value=v)
            ws.cell(row=i, column=9, value=f"=SUM(B{i}:H{i})")
        ws.cell(row=11, column=1, value="Total").font = NEGRITA
        for j in range(2, 10):
            L = get_column_letter(j)
            ws.cell(row=11, column=j, value=f"=SUM({L}3:{L}10)").font = NEGRITA
        anchos(ws, A=22)
    ws = wb.create_sheet("Resumen" if sol else "Hoja4")
    encabezados(ws, 1, ["Sucursal", "Total semana"])
    for i, s in enumerate(["Centro", "Norte", "Sur", "Total"], 2):
        ws.cell(row=i, column=1, value=s)
    anchos(ws, A=14, B=14)
    if sol:
        for i, s in enumerate(["Centro", "Norte", "Sur"], 2):
            ws.cell(row=i, column=2, value=f"={s}!I11")
        ws["B5"] = "=SUM(B2:B4)"
        for s, color in (("Centro", "1E7145"), ("Norte", "2463A6"), ("Sur", "C0392B")):
            hojas[s].sheet_properties.tabColor = color
        hojas["Centro"].freeze_panes = "B3"
        wb.move_sheet("Resumen", offset=-3)
    return wb


def totales_producto():
    t = {}
    for v in D["ventas"]:
        t[v["producto"]] = t.get(v["producto"], 0) + v["total"]
    return t


def mensual(clave, valor_clave=None):
    o = {}
    for v in D["ventas"]:
        k = v[clave]
        o.setdefault(k, [0] * 12)[int(v["fecha"][5:7]) - 1] += v["total"]
    return o


def lab06(sol):
    wb = libro("Laboratorio 6 · Gráficos para la junta", "Tres hojas: Productos (ventas anuales), Meses (ventas mensuales por categoría) y Sucursales.",
               ["Gráfico de columnas o barras de las ventas por producto (Productos!A1:B21), con título.", "Gráfico de líneas de Café caliente y Café frío por mes (Meses!A1:C13), con título.",
                "Minigráficos de línea en Sucursales!N2:N4 (datos B2:M4).", "Nivel Experto: títulos de ejes en ambos gráficos."])
    ws = wb.create_sheet("Productos")
    encabezados(ws, 1, ["Producto", "Ventas 2025"])
    tp = totales_producto()
    for i, p in enumerate(D["productos"], 2):
        ws.cell(row=i, column=1, value=p["nombre"])
        ws.cell(row=i, column=2, value=tp.get(p["nombre"], 0)).number_format = MONEDA0
    anchos(ws, A=24, B=14)
    wm = wb.create_sheet("Meses")
    encabezados(wm, 1, ["Mes"] + CATS)
    mc = mensual("categoria")
    for i, m in enumerate(MESES, 2):
        wm.cell(row=i, column=1, value=m)
        for j, c in enumerate(CATS, 2):
            wm.cell(row=i, column=j, value=mc[c][i - 2]).number_format = MONEDA0
    anchos(wm, A=8, B=14, C=12, D=10, E=12, F=10, G=12)
    wsu = wb.create_sheet("Sucursales")
    encabezados(wsu, 1, ["Sucursal"] + MESES + ["Tendencia"])
    ms = mensual("sucursal")
    for i, s in enumerate(["Centro", "Norte", "Sur"], 2):
        wsu.cell(row=i, column=1, value=s)
        for j in range(12):
            wsu.cell(row=i, column=j + 2, value=ms[s][j]).number_format = MONEDA0
    anchos(wsu, A=11, N=16)
    if sol:
        b = BarChart()
        b.type = "bar"
        b.title = "Ventas 2025 por producto"
        b.x_axis.title, b.y_axis.title = "Producto", "Ventas ($)"
        b.add_data(Reference(ws, min_col=2, min_row=1, max_row=21), titles_from_data=True)
        b.set_categories(Reference(ws, min_col=1, min_row=2, max_row=21))
        b.height, b.width = 12, 18
        ws.add_chart(b, "D2")
        lc = LineChart()
        lc.title = "Café caliente vs. café frío"
        lc.x_axis.title, lc.y_axis.title = "Mes", "Ventas ($)"
        lc.add_data(Reference(wm, min_col=2, max_col=3, min_row=1, max_row=13), titles_from_data=True)
        lc.set_categories(Reference(wm, min_col=1, min_row=2, max_row=13))
        wm.add_chart(lc, "I2")
        hoja_solucion(wb, ["Minigráficos: selecciona Sucursales!N2:N4 › Insertar › Minigráficos › Línea › Rango de datos B2:M4.", "(openpyxl no crea minigráficos: agrégalos en Excel.)"])
    return wb


def lab08(sol):
    wb = libro("Laboratorio 8 · Comisiones de noviembre", "Mismo cálculo de comisiones con las ventas de noviembre, y un resumen por sucursal.",
               ["D2:D21: comisión = ventas × tasa (F2) si las ventas alcanzan la meta (F1); si no, 0.", "F4: cuántos vendedores alcanzaron la meta (CONTAR.SI).",
                "F8:F10: comisiones de cada sucursal (SUMAR.SI).", "Nivel Experto: nombre «Meta» para F1 usado en la columna D; G8:G10 promedio de ventas por sucursal (PROMEDIO.SI)."])
    ws = wb.create_sheet("Ventas")
    encabezados(ws, 1, ["Vendedor", "Sucursal", "Ventas del mes", "Comisión"])
    for i, e in enumerate(D["empleados"], 2):
        ws.cell(row=i, column=1, value=e["nombre"])
        ws.cell(row=i, column=2, value=e["sucursal"])
        ws.cell(row=i, column=3, value=e["ventasMesLab"]).number_format = MONEDA0
        ws.cell(row=i, column=4).number_format = MONEDA
    for celda, txt in (("E1", "Meta mensual"), ("E2", "Tasa de comisión"), ("E4", "Alcanzaron la meta"), ("E7", "Sucursal"), ("F7", "Comisiones"), ("G7", "Promedio de ventas")):
        ws[celda] = txt
        ws[celda].font = NEGRITA
    ws["F1"], ws["F2"] = 10000, 0.05
    ws["F1"].number_format, ws["F2"].number_format = MONEDA0, "0%"
    for i, s in enumerate(["Centro", "Norte", "Sur"], 8):
        ws.cell(row=i, column=5, value=s)
        ws.cell(row=i, column=6).number_format = MONEDA
        ws.cell(row=i, column=7).number_format = MONEDA0
    anchos(ws, A=20, B=10, C=15, D=12, E=20, F=13, G=18)
    if sol:
        wb.defined_names["Meta"] = DefinedName("Meta", attr_text="Ventas!$F$1")
        for i in range(2, 22):
            ws.cell(row=i, column=4, value=f"=IF(C{i}>=Meta,C{i}*$F$2,0)")
        ws["F4"] = '=COUNTIF(C2:C21,">="&F1)'
        for i in range(8, 11):
            ws.cell(row=i, column=6, value=f"=SUMIF($B$2:$B$21,E{i},$D$2:$D$21)")
            ws.cell(row=i, column=7, value=f"=AVERAGEIF($B$2:$B$21,E{i},$C$2:$C$21)")
    return wb


def clientes_lab():
    rng = random.Random(909)
    usados, out = set(), []
    for i in range(30):
        while True:
            n, a = rng.choice(LD.NOMBRES), rng.choice(LD.APELLIDOS)
            if (n, a) not in usados:
                usados.add((n, a))
                break
        correo = f"{n.lower()}.{a.lower()}@correo.mx"
        for x, y in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n")):
            correo = correo.replace(x, y)
        out.append({"id": f"CL-{301 + i}", "nombre": f"{n} {a}", "telefono": 5500000000 + rng.randint(10000000, 89999999), "correo": correo,
                    "sucursal": rng.choice(["Centro", "Norte", "Sur"]), "registro": (dt.date(2024, 3, 1) + dt.timedelta(days=rng.randint(0, 500))).isoformat(), "puntos": rng.randint(0, 50) * 10})
    sucio = list(out)
    for k, idx in enumerate([2, 8, 14, 21]):
        sucio.insert(idx + 5 + k * 2, dict(out[idx]))
    return out, sucio


def lab09(sol):
    wb = libro("Laboratorio 9 · Limpiar la base de clientes", "Otra exportación de clientes con registros repetidos y el nombre completo en una sola columna.",
               ["Quita los registros duplicados.", "Separa el nombre completo: nombre en B y apellido en C (Texto en columnas).", "Ordena por Sucursal (A→Z) y luego por Apellido (A→Z).",
                "Lista desplegable en Sucursal (F2:F60) con Centro, Norte y Sur.", "Nivel Experto: columna I «Usuario» con la parte del correo antes de @ (Relleno rápido)."])
    ws = wb.create_sheet("Clientes")
    encabezados(ws, 1, ["ID", "Nombre", "Apellido", "Teléfono", "Correo", "Sucursal", "Registro", "Puntos", "Usuario"])
    limpios, sucios = clientes_lab()
    filas = sucios
    if sol:
        filas = sorted(limpios, key=lambda c: (c["sucursal"], c["nombre"].split()[1].lower()))
    for i, c in enumerate(filas, 2):
        n, a = c["nombre"].split()
        ws.cell(row=i, column=1, value=c["id"])
        ws.cell(row=i, column=2, value=n if sol else c["nombre"])
        ws.cell(row=i, column=3, value=a if sol else None)
        ws.cell(row=i, column=4, value=c["telefono"])
        ws.cell(row=i, column=5, value=c["correo"])
        ws.cell(row=i, column=6, value=c["sucursal"])
        poner_fecha(ws.cell(row=i, column=7), c["registro"])
        ws.cell(row=i, column=8, value=c["puntos"])
        if sol:
            ws.cell(row=i, column=9, value=c["correo"].split("@")[0])
    anchos(ws, A=9, B=20, C=14, D=13, E=30, F=10, G=12, H=8, I=20)
    if sol:
        dv = DataValidation(type="list", formula1='"Centro,Norte,Sur"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("F2:F60")
    return wb


def lab10(sol):
    wb = libro("Laboratorio 10 · Tabla de pedidos", "La hoja Pedidos tiene los 40 pedidos de catering (A1:I41).",
               ["Convierte A1:I41 en tabla y nómbrala «Pedidos».", "Agrega la columna calculada «Importe» = [@Cantidad]*[@Precio].", "Activa la fila de totales con la suma de Importe.",
                "Inserta una segmentación por Sucursal.", "Nivel Experto: total también en Cantidad, segmentación por Estado y otro estilo de tabla."])
    ws = wb.create_sheet("Pedidos")
    escribir_pedidos(ws, D["pedidos"])
    if sol:
        for c in ws[1]:
            c.fill = PatternFill()
        ws["J1"] = "Importe"
        for i in range(2, 42):
            ws.cell(row=i, column=10, value="=Pedidos[[#This Row],[Cantidad]]*Pedidos[[#This Row],[Precio]]").number_format = MONEDA0
        ws["A42"] = "Total"
        ws["F42"] = "=SUBTOTAL(109,Pedidos[Cantidad])"
        ws["J42"] = "=SUBTOTAL(109,Pedidos[Importe])"
        t = Table(displayName="Pedidos", ref="A1:J42", totalsRowCount=1)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
        t._initialise_columns()
        for col, cel in zip(t.tableColumns, ws[1]):
            col.name = str(cel.value)
        t.tableColumns[0].totalsRowLabel = "Total"
        t.tableColumns[5].totalsRowFunction = "sum"
        t.tableColumns[9].totalsRowFunction = "sum"
        t.tableColumns[9].calculatedColumnFormula = TableFormula(attr_text="Pedidos[[#This Row],[Cantidad]]*Pedidos[[#This Row],[Precio]]")
        ws.add_table(t)
        anchos(ws, J=12)
        hoja_solucion(wb, ["Segmentación: con una celda de la tabla › Diseño de tabla › Insertar segmentación de datos › Sucursal (y Estado).", "(openpyxl no crea segmentaciones: agrégalas en Excel.)"])
    return wb


def inventario_norte():
    rng = random.Random(1111)
    inv = [dict(x) for x in D["inventario"]]
    for x in inv:
        x["existencia"] = max(1, round(x["existencia"] * rng.uniform(0.7, 1.3)))
    return inv


def lab11(sol):
    wb = libro("Laboratorio 11 · Semáforo del inventario de Norte", "Mismo inventario con las existencias de Norte.",
               ["C2:C16: relleno rojo cuando la existencia sea menor que el mínimo de su fila (regla «Es menor que» con =$D2).", "H2:H16: conjunto de iconos (semáforo).",
                "F2:F16: barras de datos.", "A1:H1: estilo de celda Encabezado 1.", "Nivel Experto: amarillo entre el mínimo y 1.3 veces el mínimo; cambia el tema del libro."])
    ws = wb.create_sheet("Inventario")
    escribir_inventario(ws, inventario_norte())
    if sol:
        for c in ws[1]:
            c.style = "Headline 1"
        ws.conditional_formatting.add("C2:C16", CellIsRule(operator="lessThan", formula=["$D2"], fill=ROJO, font=Font(color="9C0006")))
        ws.conditional_formatting.add("C2:C16", CellIsRule(operator="between", formula=["$D2", "$D2*1.3"], fill=AMARILLO, font=Font(color="9C5700")))
        ws.conditional_formatting.add("H2:H16", IconSetRule("3TrafficLights1", "percent", [0, 33, 67]))
        ws.conditional_formatting.add("F2:F16", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    return wb


def lideres():
    t = {}
    for v in D["ventas"]:
        q = (int(v["fecha"][5:7]) - 1) // 3 + 1
        t.setdefault(q, {}).setdefault(v["sucursal"], 0)
        t[q][v["sucursal"]] += v["total"]
    return t, [max(t[q], key=t[q].get) for q in (1, 2, 3, 4)]


# ================================================================ versión 2
def lab_inicio(sol):
    wb = libro("Laboratorio 1 · Mi primer libro en Excel", "Recorre la pantalla de Excel haciendo cambios sencillos que el verificador puede revisar.",
               ["Cambia el nombre de «Hoja1» a «Mi primera hoja» (doble clic en la pestaña).",
                "En A1 escribe «Mi primer libro en La Taza». Con el cuadro de nombres ve a D10 y escribe tu nombre.",
                "Ajusta el zoom de la hoja a 120 % (barra de estado o Vista › Zoom).",
                "Archivo › Información › Propiedades: escribe el título «Mi primer libro».",
                "Configura la Autorrecuperación cada 5 minutos (Archivo › Opciones › Guardar) y personaliza tu barra de acceso rápido.",
                "Nivel Experto: deja la hoja en la vista Diseño de página y agrega «curso» a las palabras clave (Etiquetas)."],
               ["Nombre de la hoja, A1 y D10.", "Zoom de la hoja y título del documento.", "Experto: vista Diseño de página y palabra clave."])
    ws = wb.create_sheet("Mi primera hoja" if sol else "Hoja1")
    if sol:
        ws["A1"] = "Mi primer libro en La Taza"
        ws["D10"] = "Ana Torres"
        ws.sheet_view.zoomScale = 120
        ws.sheet_view.view = "pageLayout"
        wb.properties.title = "Mi primer libro"
        wb.properties.keywords = "curso; La Taza"
        hoja_solucion(wb, ["La Autorrecuperación y la barra de acceso rápido son opciones de Excel, no del libro: se revisan en el simulador (reto del módulo 1).",
                           "Título y palabras clave: Archivo › Información › Propiedades (o Propiedades avanzadas › Resumen)."])
    wb.active = 1
    return wb


def mejor_mes_norte():
    m = [0] * 12
    for v in D["ventas"]:
        if v["sucursal"] == "Norte":
            m[int(v["fecha"][5:7]) - 1] += v["total"]
    return D["meses"][m.index(max(m))]


def lider_anual():
    t = {}
    for v in D["ventas"]:
        t[v["sucursal"]] = t.get(v["sucursal"], 0) + v["total"]
    return max(t, key=t.get), t


def lab_dinamicas(sol):
    wb = libro("Laboratorio 12 · Tablas dinámicas y tablero", "La hoja Ventas tiene las 720 ventas de 2025. La hoja Respuestas es para tus conclusiones.",
               ["Inserta una tabla dinámica: Fecha agrupada por trimestres en filas, Sucursal en columnas y suma de Total en valores.",
                "En Respuestas escribe la sucursal que más vendió en cada trimestre (B2:B5), el mes con más ventas de Norte (B7) y la sucursal que más vendió en el año (B8).",
                "Tablero: otra tabla dinámica por meses, un gráfico dinámico y una segmentación por Sucursal conectada.",
                "Nivel Experto: valores como % del total en alguna tabla dinámica y una escala de tiempo por Fecha."],
               ["Tabla dinámica por trimestre y sucursal.", "Respuestas B2:B5, B7 y B8.", "Gráfico dinámico y segmentación conectada."])
    hoja_ventas(wb)
    ws = wb.create_sheet("Respuestas")
    encabezados(ws, 1, ["Trimestre", "Sucursal que más vendió"])
    t, lid = lideres()
    for i in range(4):
        ws.cell(row=i + 2, column=1, value=f"T{i + 1}")
    ws["A7"], ws["A8"] = "Mes con más ventas de Norte", "Sucursal que más vendió en el año"
    ws["A7"].font = ws["A8"].font = NEGRITA
    for d in ("B2", "B3", "B4", "B5", "B7", "B8"):
        ws[d].fill = RELLENO_ENTRADA
    anchos(ws, A=34, B=26)
    if sol:
        for i in range(4):
            ws.cell(row=i + 2, column=2, value=lid[i])
        ws["B7"] = mejor_mes_norte()
        lider, tot = lider_anual()
        ws["B8"] = lider
        s = hoja_solucion(wb, ["1. Clic en Ventas › Insertar › Tabla dinámica › Nueva hoja.", "2. Fecha a Filas; clic derecho en una fecha › Agrupar › solo Trimestres.",
                               "3. Sucursal a Columnas; Total a Valores (Suma).", "4. Tablero: Fecha (Meses) en filas, Total en valores; Gráfico dinámico › Líneas; Insertar segmentación › Sucursal.",
                               "5. Experto: Mostrar valores como › % del total general; Insertar escala de tiempo › Fecha.",
                               "openpyxl no crea tablas dinámicas: estos son los resultados esperados.", "Totales por trimestre:"])
        encabezados(s, 11, ["Trimestre", "Centro", "Norte", "Sur"])
        for q in range(1, 5):
            s.cell(row=11 + q, column=1, value=f"T{q}")
            for j, su in enumerate(["Centro", "Norte", "Sur"], 2):
                s.cell(row=11 + q, column=j, value=t[q][su]).number_format = MONEDA0
        s.cell(row=16, column=1, value="Año").font = NEGRITA
        for j, su in enumerate(["Centro", "Norte", "Sur"], 2):
            s.cell(row=16, column=j, value=tot[su]).number_format = MONEDA0
    return wb


# ---------------------------------------------------- gestión de pedidos (módulo 13)
CLIENTES = D["clientesCorp"]
PRODUCTOS = D["productos"]


def hoja_catalogo(wb):
    ws = wb.create_sheet("Catálogo")
    encabezados(ws, 1, ["Producto", "Precio", "Categoría"])
    for i, p in enumerate(PRODUCTOS, 2):
        ws.cell(row=i, column=1, value=p["nombre"])
        ws.cell(row=i, column=2, value=p["precio"]).number_format = MONEDA0
        ws.cell(row=i, column=3, value=p["categoria"])
    anchos(ws, A=22, B=9, C=15)
    return ws


def hoja_clientes(wb):
    ws = wb.create_sheet("Clientes")
    encabezados(ws, 1, ["Cliente", "Contacto", "Correo", "Teléfono"])
    for i, c in enumerate(CLIENTES, 2):
        for j, k in enumerate(["cliente", "contacto", "correo", "telefono"], 1):
            ws.cell(row=i, column=j, value=c[k])
    anchos(ws, A=24, B=22, C=26, D=13)
    return ws


def validaciones_pedidos(ws, hasta):
    dv = DataValidation(type="list", formula1="Clientes!$A$2:$A$11", allow_blank=True, showErrorMessage=True, errorTitle="Cliente", error="Elige un cliente de la lista.")
    dv.add(f"C2:C{hasta}")
    dp = DataValidation(type="list", formula1="'Catálogo'!$A$2:$A$21", allow_blank=True, showErrorMessage=True, errorTitle="Producto", error="Elige un producto del catálogo.")
    dp.add(f"D2:D{hasta}")
    dc = DataValidation(type="whole", operator="between", formula1="1", formula2="200", allow_blank=True, showErrorMessage=True,
                        showInputMessage=True, promptTitle="Cantidad", prompt="Entero de 1 a 200", errorTitle="Cantidad", error="Escribe un número entero de 1 a 200.")
    dc.add(f"E2:E{hasta}")
    for v in (dv, dp, dc):
        ws.add_data_validation(v)


def hoja_pedidos_gestion(wb, sol, filas=None):
    ws = wb.create_sheet("Pedidos")
    encabezados(ws, 1, ["Folio", "Fecha", "Cliente", "Producto", "Cantidad", "Precio", "Importe", "Entrega"])
    ped = filas or D["pedidos"]
    for i, p in enumerate(ped, 2):
        ws.cell(row=i, column=1, value=p["folio"])
        poner_fecha(ws.cell(row=i, column=2), p["fecha"])
        ws.cell(row=i, column=3, value=p["cliente"])
        ws.cell(row=i, column=4, value=p["producto"])
        ws.cell(row=i, column=5, value=p["cantidad"])
        poner_fecha(ws.cell(row=i, column=8), p["compromiso"])
        ws.cell(row=i, column=6).number_format = MONEDA
        ws.cell(row=i, column=7).number_format = MONEDA
        if sol:
            ws.cell(row=i, column=6, value=f"=_xlfn.XLOOKUP(D{i},Catálogo!$A$2:$A$21,Catálogo!$B$2:$B$21)")
            ws.cell(row=i, column=7, value=f"=E{i}*F{i}")
    ws.freeze_panes = "A2"
    anchos(ws, A=8, B=11, C=24, D=22, E=10, F=10, G=12, H=11)
    if sol:
        validaciones_pedidos(ws, 60)
    return ws


def hoja_consulta(wb, sol, cliente="Notaría 14"):
    ws = wb.create_sheet("Consulta")
    ws["B1"] = "Consulta de clientes"
    ws["B1"].font = TITULO
    ws["B2"], ws["C2"] = "Cliente:", cliente
    ws["C2"].fill = RELLENO_ENTRADA
    ws["C2"].font = NEGRITA
    ws["B4"], ws["B5"] = "Número de pedidos", "Total comprado"
    ws["B4"].font = ws["B5"].font = NEGRITA
    ws["C5"].number_format = MONEDA
    anchos(ws, A=3, B=22, C=26)
    if sol:
        ws["C4"] = "=COUNTIF(Pedidos!C2:C60,C2)"
        ws["C5"] = "=SUMIF(Pedidos!C2:C60,C2,Pedidos!G2:G60)"
    return ws


def lab_gestion(sol):
    wb = libro("Laboratorio 13 · Registro de pedidos y cartas de confirmación",
               "Parte 1 en Excel: base de pedidos con captura validada, búsqueda y consulta. Parte 2 en Word: cartas con combinar correspondencia (lab13-carta.docx).",
               ["Pedidos: lista desplegable de clientes en C2:C60 (origen =Clientes!$A$2:$A$11) y de productos en D2:D60 (origen =Catálogo!$A$2:$A$21); en E2:E60 solo enteros de 1 a 200.",
                "Pedidos: Precio (F) con BUSCARX desde el Catálogo e Importe (G) = Cantidad × Precio, para todos los pedidos.",
                "Consulta: C4 = número de pedidos del cliente de C2 (CONTAR.SI) y C5 = total comprado (SUMAR.SI).",
                "Guarda y cierra el libro. En Word abre lab13-carta.docx: Correspondencia › Iniciar combinación › Cartas › Seleccionar destinatarios › Usar una lista existente (este libro, hoja Pedidos).",
                "Inserta los campos combinados donde dice [Folio], [Cliente], [Cantidad], [Producto], [Importe] y [Entrega]. Revisa la vista previa y guarda la carta.",
                "Sube el libro y la carta juntos.",
                "Nivel Experto: macro que limpie Consulta!C2 asignada a un botón (guarda como .xlsm) y formato de moneda en el campo Importe con \\# \"$#,##0.00\"."],
               ["Validaciones de captura.", "Precio e importe con fórmulas.", "Consulta que cambia con el cliente de C2.", "Carta con los campos combinados."])
    hoja_pedidos_gestion(wb, sol)
    hoja_catalogo(wb)
    hoja_clientes(wb)
    hoja_consulta(wb, sol)
    if sol:
        hoja_solucion(wb, ["Validación: Datos › Validación de datos › Permitir Lista › Origen =Clientes!$A$2:$A$11 (y =Catálogo!$A$2:$A$21); Número entero entre 1 y 200.",
                           "F2: =BUSCARX(D2,Catálogo!$A$2:$A$21,Catálogo!$B$2:$B$21) · G2: =E2*F2 (copiar hasta la última fila).",
                           "Consulta: C4 =CONTAR.SI(Pedidos!C2:C60,C2) · C5 =SUMAR.SI(Pedidos!C2:C60,C2,Pedidos!G2:G60).",
                           "Experto (macro): Programador › Grabar macro «LimpiarConsulta» › seleccionar Consulta!C2 › Supr › Detener; botón con la macro; guardar como .xlsm.",
                           "Sub LimpiarConsulta()", '    Range("C2").Select', "    Selection.ClearContents", "End Sub",
                           "La carta resuelta (lab13-carta-solucion.docx) muestra los campos combinados y el modificador de moneda en Importe."])
    return wb


# ---------------------------------------------------- documento de Word mínimo
def _p(texto="", negrita=False, tam=None, centrado=False, color=None):
    """Párrafo; en el texto, {Campo} o {Campo|modificador} se convierte en campo combinado."""
    ppr = '<w:pPr><w:jc w:val="center"/></w:pPr>' if centrado else ""
    rpr = ""
    if negrita or tam or color:
        rpr = "<w:rPr>" + ("<w:b/>" if negrita else "") + (f'<w:color w:val="{color}"/>' if color else "") + (f'<w:sz w:val="{tam * 2}"/>' if tam else "") + "</w:rPr>"
    runs = []
    for parte in re.split(r"(\{[^}]+\})", texto):
        if not parte:
            continue
        m = re.fullmatch(r"\{([^}|]+)(?:\|([^}]*))?\}", parte)
        if m:
            campo, mod = m.group(1), m.group(2)
            instr = f" MERGEFIELD {campo} " + (mod + " " if mod else "")
            runs.append(f'<w:fldSimple w:instr="{escape(instr, {chr(34): "&quot;"})}"><w:r>{rpr}<w:t>«{escape(campo)}»</w:t></w:r></w:fldSimple>')
        else:
            runs.append(f'<w:r>{rpr}<w:t xml:space="preserve">{escape(parte)}</w:t></w:r>')
    return f"<w:p>{ppr}{''.join(runs)}</w:p>"


def docx(ruta, parrafos, titulo):
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    documento = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>{"".join(parrafos)}'
                 '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1418" w:right="1418" w:bottom="1418" w:left="1418" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr>'
                 "</w:body></w:document>")
    estilos = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {ns}><w:docDefaults><w:rPrDefault><w:rPr>'
               '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="Calibri" w:cs="Calibri"/><w:sz w:val="22"/><w:lang w:val="es-MX"/></w:rPr></w:rPrDefault>'
               '<w:pPrDefault><w:pPr><w:spacing w:after="160" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
               '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style></w:styles>')
    ajustes = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:settings {ns}><w:defaultTabStop w:val="708"/><w:compat/></w:settings>'
    tipos = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
             '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
             '<Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>'
             '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/></Relationships>')
    rels_doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" Target="settings.xml"/></Relationships>')
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{escape(titulo)}</dc:title><dc:creator>Cafetería La Taza</dc:creator></cp:coreProperties>')
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", tipos)
        z.writestr("_rels/.rels", rels)
        z.writestr("docProps/core.xml", core)
        z.writestr("word/document.xml", documento)
        z.writestr("word/_rels/document.xml.rels", rels_doc)
        z.writestr("word/styles.xml", estilos)
        z.writestr("word/settings.xml", ajustes)


def carta(sol):
    if sol:
        c = lambda campo, mod=None: "{" + campo + ("|" + mod if mod else "") + "}"
    else:
        c = lambda campo, mod=None: "[" + campo + "]"
    return [
        _p("Cafetería La Taza", negrita=True, tam=20, color="5B3A29"),
        _p("Pedidos para empresas · catering@lataza.example", color="7A6A5E"),
        _p(""),
        _p("Confirmación de pedido", negrita=True, tam=14),
        _p("Pedido número: " + c("Folio")),
        _p("Cliente: " + c("Cliente")),
        _p(""),
        _p("Estimado cliente:"),
        _p("Le confirmamos su pedido de " + c("Cantidad") + " piezas de " + c("Producto") + ", por un importe de "
           + c("Importe", '\\# "$#,##0.00"') + ". La entrega está programada para el " + c("Entrega", '\\@ "dd/MM/yyyy"') + "."),
        _p("Si necesita algún cambio, responda a este correo antes de la fecha de entrega."),
        _p("¡Gracias por su preferencia!"),
        _p(""),
        _p("Marta Ríos", negrita=True),
        _p("Gerente general · Cafetería La Taza"),
    ]


# ================================================================ integradores
def int_basico(sol):
    wb = libro("Reto integrador básico · Reporte semanal de ventas", "Reporte semanal de la sucursal Sur: cálculos, formato, gráfico, logo y resumen.",
               ["Semana: J4:J11 unidades de la semana (suma de Lun a Dom); K4:K11 ventas (precio × unidades); J12 y K12 totales; N4 promedio de ventas por producto; N5 venta mayor.",
                "Formato: título destacado, encabezados A3:K3 en negrita con relleno, precio y ventas en moneda, bordes en A3:K12.",
                "Gráfico de columnas o barras de las ventas por producto, con título.",
                "Inserta el logo de La Taza (logo-la-taza.png) en la hoja Semana.",
                "Agrega una hoja llamada «Resumen» con el total de la semana en B2 (fórmula que tome Semana!K12).",
                "Escribe el título del documento «Reporte semanal Sur» (Archivo › Información › Propiedades)."])
    ws = wb.create_sheet("Semana")
    ws["A1"] = "Reporte semanal · Sucursal Sur · Semana 42"
    cab = ["Producto", "Precio"] + DIAS + ["Unidades", "Ventas"]
    for j, v in enumerate(cab, 1):
        ws.cell(row=3, column=j, value=v)
    for i, p in enumerate(D["semanaSucursal"]["Sur"], 4):
        ws.cell(row=i, column=1, value=p["producto"])
        ws.cell(row=i, column=2, value=p["precio"])
        for j, v in enumerate(p["cantidades"], 3):
            ws.cell(row=i, column=j, value=v)
    ws["A12"] = "Total"
    ws["M4"], ws["M5"] = "Promedio por producto", "Venta mayor"
    anchos(ws, A=22, M=22)
    if sol:
        ws["A1"].font = TITULO
        for i in range(4, 12):
            ws.cell(row=i, column=10, value=f"=SUM(C{i}:I{i})")
            ws.cell(row=i, column=11, value=f"=B{i}*J{i}")
        ws["J12"], ws["K12"] = "=SUM(J4:J11)", "=SUM(K4:K11)"
        ws["N4"], ws["N5"] = "=AVERAGE(K4:K11)", "=MAX(K4:K11)"
        for c in ws[3]:
            if c.column <= 11:
                c.font, c.fill = NEGRITA, RELLENO_ENC
        for r in range(4, 13):
            ws.cell(row=r, column=2).number_format = MONEDA
            ws.cell(row=r, column=11).number_format = MONEDA
        ws["N4"].number_format = ws["N5"].number_format = MONEDA
        bordes(ws, "A3:K12")
        ch = BarChart()
        ch.title = "Ventas de la semana por producto"
        ch.add_data(Reference(ws, min_col=11, min_row=3, max_row=11), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=11))
        ws.add_chart(ch, "A15")
        img = Image(str(RAIZ / "assets" / "img" / "logo-la-taza.png"))
        img.width, img.height = 60, 60
        ws.add_image(img, "P1")
        wr = wb.create_sheet("Resumen")
        wr["A2"], wr["B2"] = "Total de la semana", "=Semana!K12"
        wr["B2"].number_format = MONEDA
        anchos(wr, A=22, B=14)
        wb.properties.title = "Reporte semanal Sur"
    return wb


def int_intermedio(sol):
    wb = libro("Reto integrador intermedio · Base de pedidos confiable", "Limpia la base de pedidos y calcula las comisiones del mes.",
               ["Pedidos: quita los duplicados.", "Convierte los pedidos en la tabla «Pedidos» con la columna calculada Importe y la fila de totales.",
                "Formato condicional para los pedidos con Estado «Pendiente».", "Lista desplegable en la columna Estado (I2:I41) con Pendiente y Entregado.",
                "Comisiones: D2:D21 (meta F1, tasa F2) y F4 (cuántos alcanzaron la meta)."])
    ws = wb.create_sheet("Pedidos")
    ped = list(D["pedidos"])
    if not sol:
        for k, idx in enumerate([4, 11, 19, 26, 33]):
            ped.insert(idx + k, dict(D["pedidos"][idx]))
    escribir_pedidos(ws, ped)
    wc = wb.create_sheet("Comisiones")
    encabezados(wc, 1, ["Vendedor", "Sucursal", "Ventas del mes", "Comisión"])
    rng = random.Random(7070)
    for i, e in enumerate(D["empleados"], 2):
        wc.cell(row=i, column=1, value=e["nombre"])
        wc.cell(row=i, column=2, value=e["sucursal"])
        wc.cell(row=i, column=3, value=round(e["ventasMes"] * rng.uniform(0.9, 1.1))).number_format = MONEDA0
        wc.cell(row=i, column=4).number_format = MONEDA
    for celda, txt in (("E1", "Meta mensual"), ("E2", "Tasa de comisión"), ("E4", "Alcanzaron la meta")):
        wc[celda] = txt
        wc[celda].font = NEGRITA
    wc["F1"], wc["F2"] = 10000, 0.05
    wc["F1"].number_format, wc["F2"].number_format = MONEDA0, "0%"
    anchos(wc, A=20, C=15, E=20)
    if sol:
        for c in ws[1]:
            c.fill = PatternFill()
        ws["J1"] = "Importe"
        for i in range(2, 42):
            ws.cell(row=i, column=10, value="=Pedidos[[#This Row],[Cantidad]]*Pedidos[[#This Row],[Precio]]").number_format = MONEDA0
        ws["A42"], ws["J42"] = "Total", "=SUBTOTAL(109,Pedidos[Importe])"
        t = Table(displayName="Pedidos", ref="A1:J42", totalsRowCount=1)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        t._initialise_columns()
        for col, cel in zip(t.tableColumns, ws[1]):
            col.name = str(cel.value)
        t.tableColumns[0].totalsRowLabel = "Total"
        t.tableColumns[9].totalsRowFunction = "sum"
        t.tableColumns[9].calculatedColumnFormula = TableFormula(attr_text="Pedidos[[#This Row],[Cantidad]]*Pedidos[[#This Row],[Precio]]")
        ws.add_table(t)
        ws.conditional_formatting.add("I2:I41", CellIsRule(operator="equal", formula=['"Pendiente"'], fill=AMARILLO))
        dv = DataValidation(type="list", formula1='"Pendiente,Entregado"', allow_blank=True)
        dv.add("I2:I41")
        ws.add_data_validation(dv)
        for i in range(2, 22):
            wc.cell(row=i, column=4, value=f"=IF(C{i}>=$F$1,C{i}*$F$2,0)")
        wc["F4"] = '=COUNTIF(C2:C21,">="&F1)'
    return wb


def int_avanzado(sol):
    wb = libro("Reto integrador avanzado · Tablero y gestión automatizada",
               "Tablero de ventas con tablas dinámicas y un registro de pedidos con captura validada, búsqueda, consulta y una macro. Entrega en .xlsm.",
               ["Hoja nueva «Tablero»: tabla dinámica con la suma de Total por mes, gráfico dinámico y segmentación por Sucursal.",
                "Pedidos: listas desplegables de clientes (C2:C60) y productos (D2:D60) y cantidad entera de 1 a 200 (E2:E60).",
                "Pedidos: Precio (F) con BUSCARX desde el Catálogo e Importe (G) = Cantidad × Precio.",
                "Consulta: C4 número de pedidos y C5 total comprado del cliente de C2.",
                "Una macro que limpie Consulta!C2, asignada a un botón. Guarda como .xlsm."])
    hoja_ventas(wb)
    hoja_pedidos_gestion(wb, sol, D["pedidos"][:30])
    hoja_catalogo(wb)
    hoja_clientes(wb)
    hoja_consulta(wb, sol, "Hospital San Rafael")
    if sol:
        hoja_solucion(wb, ["Tablero: Insertar › Tabla dinámica (Fecha por meses, Suma de Total), Gráfico dinámico, Segmentación por Sucursal.",
                           "Macro: Programador › Grabar macro «LimpiarConsulta» (seleccionar Consulta!C2, Supr); botón con la macro; guardar como .xlsm.",
                           "openpyxl no crea tablas dinámicas ni guarda VBA: esta solución incluye las validaciones, las fórmulas y la consulta."])
    return wb


def main():
    libro_base()
    for nombre, f in [("lab01-mi-primer-libro", lab_inicio), ("lab02-lista-precios", lab01), ("lab03-corte-caja", lab02), ("lab04-reporte-semanal", lab03),
                      ("lab05-menu", lab04), ("lab06-sucursales", lab05), ("lab07-graficos", lab06), ("lab08-comisiones", lab08),
                      ("lab09-clientes", lab09), ("lab10-pedidos", lab10), ("lab11-inventario", lab11), ("lab12-tablas-dinamicas", lab_dinamicas),
                      ("lab13-gestion-pedidos", lab_gestion)]:
        lab(nombre, f)
    docx(RAIZ / "laboratorios" / "lab13-carta.docx", carta(False), "Confirmación de pedido")
    docx(PRIVADO / "soluciones" / "laboratorios" / "lab13-carta-solucion.docx", carta(True), "Confirmación de pedido (solución)")
    integrador("int-basico", int_basico)
    integrador("int-intermedio", int_intermedio)
    integrador("int-avanzado", int_avanzado)
    print("libros generados · soluciones en", PRIVADO / "soluciones")


if __name__ == "__main__":
    main()
